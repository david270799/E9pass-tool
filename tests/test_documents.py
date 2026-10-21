"""Догрузка документов: файл без расширения, пароль из даты рождения, повторное определение типа."""

import pikepdf
import pymupdf
import pytest

from e9pass import services
from e9pass.db import Database


class NoOcr:
    def lines(self, img):
        return []

    def read_image_lines(self, path):
        return []


def make_pdf(path, text="hello", password=""):
    doc = pymupdf.open()
    doc.new_page().insert_text((60, 80), text)
    doc.save(path)
    doc.close()
    if password:
        with pikepdf.open(path, allow_overwriting_input=True) as pdf:
            pdf.save(path, encryption=pikepdf.Encryption(user=password, owner=password + "x"))


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("E9PASS_HOME", str(tmp_path / "home"))
    db = Database(tmp_path / "db.sqlite3")
    folder = tmp_path / "home" / "periods" / "2026-08" / "client"
    folder.mkdir(parents=True)
    cid = db.add_client("2026-08", full_name="TEST", arc_number="950101-5123456", birth_date="950101",
                        folder_path=str(folder))
    return db, cid, tmp_path


def test_file_without_extension_is_unlocked_and_classified_after_rule_change(env):
    db, cid, tmp = env
    rule = next(r for r in db.rules() if r.code == "insurance")
    rule.filename_patterns = [p for p in rule.filename_patterns if "nhis" not in p]  # будто маски еще нет
    db.save_rule(rule)
    src = tmp / "nhis-20260730162847448hzr9ewQYqbX1b0T+"   # как скачивается с сайта — без .pdf
    make_pdf(src, "insurance", password="950101")
    services.add_documents(db, cid, [src], engine=NoOcr())
    doc = db.documents(cid)[0]
    assert doc.path.suffix == ".pdf"
    assert doc.password_status == "unlocked"
    assert doc.doc_type == ""                       # маски *nhis* еще нет

    rule = next(r for r in db.rules() if r.code == "insurance")
    rule.filename_patterns.append("*nhis*")
    db.save_rule(rule)
    assert services.recheck_all_unknown(db) == 1
    assert db.documents(cid)[0].doc_type == "insurance"


def test_old_document_without_extension_fixed_on_master_build(env):
    db, cid, tmp = env
    folder = db.client(cid).folder
    raw = folder / "nhis-abc"                       # добавлен старой версией, без расширения и с паролем
    make_pdf(raw, password="19950101")
    doc_id = db.add_document(cid, "insurance", raw, original_name="nhis-abc", password_status="locked")
    res = services.build_master(db, cid)
    doc = next(d for d in db.documents(cid) if d.id == doc_id)
    assert doc.path.suffix == ".pdf" and doc.password_status == "unlocked"
    assert res.path and "Страховка" in res.included and not res.locked


def test_wrong_password_reported(env):
    db, cid, tmp = env
    src = tmp / "secret.pdf"
    make_pdf(src, password="0000")
    services.add_documents(db, cid, [src], engine=NoOcr())
    rule_doc = db.documents(cid)[0]
    db.update_document(rule_doc.id, doc_type="contract")
    res = services.build_master(db, cid)
    assert res.locked == ["secret.pdf"]
    assert services.unlock_with_password(db, rule_doc.id, cid, "0000")


def test_mask_without_star_means_contains(env):
    db, cid, tmp = env
    rule = next(r for r in db.rules() if r.code == "insurance")
    rule.filename_patterns = ["NHIS"]
    code, _ = services.classify_file(tmp / "x.pdf", [rule], None, original_name="nhis-2026.pdf")
    assert code == "insurance"


def test_delete_client_removes_folder(env):
    db, cid, tmp = env
    folder = db.client(cid).folder
    (folder / "a.txt").write_text("x")
    services.delete_client_everything(db, cid)
    assert db.client(cid) is None and not folder.exists()


def test_registry_has_upper_and_lower_name(env, tmp_path):
    from openpyxl import load_workbook
    db, cid, tmp = env
    db.update_client(cid, full_name="EGAMBERDIEV AZAM- AT PULATOVICH")
    out = services.export_registry(db.clients("2026-08"), tmp / "r.xlsx")
    rows = list(load_workbook(out).active.iter_rows(values_only=True))
    assert rows[0][:3] == ("№", "ФИО (заглавными)", "ФИО (строчными)")
    assert rows[1][1:3] == ("EGAMBERDIEV AZAMAT PULATOVICH", "egamberdiev azamat pulatovich")


def test_registry_contacts_roundtrip(env, tmp_path):
    from openpyxl import Workbook
    db, cid, tmp = env
    other = db.add_client("2026-08", full_name="TURSUNOVA MADINA", arc_number="000315-8234567",
                          folder_path=str(tmp / "x"))
    wb = Workbook()
    ws = wb.active
    ws.append(["№", "ФИО", "Номер ARC", "Телефон", "E-mail"])
    ws.append([1, "TEST", "950101-5123456", "01012345678", "Test@Mail.com"])
    ws.append([2, "tursunova madina 010-9876-5432 madina@gmail.com"])   # все в одной ячейке, без ARC
    ws.append([3, "UNKNOWN PERSON", "", "010-1111-2222", ""])
    wb.save(tmp / "contacts.xlsx")
    res = services.import_registry_contacts(db, tmp / "contacts.xlsx", period="2026-08")
    assert db.client(cid).phone == "010-1234-5678" and db.client(cid).email == "test@mail.com"
    assert db.client(other).phone == "010-9876-5432" and db.client(other).email == "madina@gmail.com"
    assert res.unmatched == ["UNKNOWN PERSON"]


def test_registry_reimport_updates_changed_address(env, tmp_path):
    from openpyxl import load_workbook
    db, cid, tmp = env
    db.update_client(cid, registered_address="경기도 화성시 항남음 발안로 45")   # OCR с ошибкой
    out = services.export_registry(db.clients("2026-08"), tmp / "r.xlsx")
    wb = load_workbook(out)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    ws.cell(2, hdr.index("Адрес регистрации") + 1, "경기도 화성시 향남읍 발안로 45, 101동 202호")
    ws.cell(2, hdr.index("Телефон") + 1, "010-5555-6666")
    wb.save(tmp / "r2.xlsx")
    res = services.import_registry_contacts(db, tmp / "r2.xlsx", period="2026-08")
    c = db.client(cid)
    assert c.registered_address == "경기도 화성시 향남읍 발안로 45, 101동 202호"
    assert c.phone == "010-5555-6666" and res.addresses_changed == 1
    # повторный импорт того же файла ничего не меняет
    assert services.import_registry_contacts(db, tmp / "r2.xlsx", period="2026-08").addresses_changed == 0


class FakeOcr:
    """OCR-заглушка: «лицо ID» для файлов с front в имени."""
    def read_image_lines(self, path):
        if "front" in str(path):
            return ["외국인등록증 RESIDENCE CARD", "900101-5123452", "성명 Name", "KARIMOV JASUR",
                    "국가/지역 UZBEKISTAN", "발급일자 2023.05.12", "CHIEF, SUWON OFFICE"]
        return []

    def read_front(self, path):
        from e9pass.ocr.front import parse_card
        lines = self.read_image_lines(path)
        return parse_card(lines), lines

    def lines(self, img):
        return []

    def lines_best_rotation(self, img, upscale=False):
        return []


def test_manual_client_filled_from_dropped_photos(tmp_path, monkeypatch):
    monkeypatch.setenv("E9PASS_HOME", str(tmp_path / "home"))
    from PIL import Image
    db = Database(tmp_path / "db.sqlite3")
    cid = services.create_manual_client(db, "2026-10")
    c = db.client(cid)
    assert c.folder.exists() and not c.full_name and "нет фото лицевой стороны ID" in c.warnings
    photo = tmp_path / "front_photo.jpg"
    Image.new("RGB", (800, 500), "white").save(photo)
    monkeypatch.setattr(services, "make_id_a4", lambda db_, cid_, path, *a, **k: None)
    services.add_documents(db, cid, [photo], engine=FakeOcr())
    c = db.client(cid)
    assert c.full_name == "KARIMOV JASUR" and c.arc_number == "900101-5123452"
    assert db.documents(cid)[0].doc_type == "id_front"
    assert "нет фото лицевой стороны ID" not in c.warnings
    # ручная правка не затирается следующим фото
    services.set_client_fields(db, cid, full_name="KARIMOV JASUR BAKHTIYOR UGLI")
    photo2 = tmp_path / "front_photo2.jpg"
    Image.new("RGB", (800, 500), "white").save(photo2)
    services.add_documents(db, cid, [photo2], engine=FakeOcr())
    assert db.client(cid).full_name == "KARIMOV JASUR BAKHTIYOR UGLI"


def test_client_zip_is_flat_without_contract_and_scans(env, monkeypatch):
    import zipfile
    db, cid, tmp = env
    folder = db.client(cid).folder
    (folder / "originals").mkdir(exist_ok=True)
    (folder / "scans").mkdir(exist_ok=True)
    files = {}
    for name, kind, scan in [("front.jpg", "id_front", 0), ("TEST.pdf", "contract", 0), ("nhis-1.pdf", "insurance", 0),
                             ("ID_A4.pdf", "id_a4", 1), ("ID_back.png", "id_back", 1)]:
        p = folder / ("scans" if scan else "originals") / name
        p.write_bytes(b"x")
        db.add_document(cid, kind, p, original_name=name, is_processed_scan=scan)
        files[name] = p
    master = folder / "MASTER.pdf"
    master.write_bytes(b"m")
    db.update_client(cid, master_pdf=str(master))
    out = services.export_client_zip(db, cid, tmp)
    names = sorted(zipfile.ZipFile(out).namelist())
    assert names == ["MASTER.pdf", "front.jpg", "nhis-1.pdf"]
