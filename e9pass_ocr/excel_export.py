"""Сохранение результатов в Excel (.xlsx)."""

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (ключ, заголовок, ширина)
COLUMNS = [
    ("file", "Файл", 22),
    ("reg_no", "Рег. номер (외국인등록번호)", 20),
    ("name", "ФИО (성명)", 28),
    ("country", "Страна (국가/지역)", 16),
    ("country_ru", "Страна (рус.)", 16),
    ("sex", "Пол", 6),
    ("birth_date", "Дата рождения", 14),
    ("visa", "Виза", 8),
    ("visa_full", "Статус (체류자격)", 20),
    ("issue_date", "Дата выдачи (발급일자)", 16),
    ("issuer", "Орган выдачи", 26),
    ("warnings", "Проверить", 34),
]

_WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def save_xlsx(rows: Iterable[dict], path) -> Path:
    """rows — словари с ключами из COLUMNS; warnings — строка или список."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ID-карты"
    ws.append([title for _, title, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for col, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    for row in rows:
        values = []
        for key, _, _ in COLUMNS:
            v = row.get(key, "")
            if isinstance(v, (list, tuple)):
                v = "; ".join(v)
            values.append(v)
        ws.append(values)
        r = ws.max_row
        for cell in ws[r]:
            cell.number_format = "@"  # всё как текст, чтобы Excel не портил номера и даты
        if values[-1]:
            for cell in ws[r]:
                cell.fill = _WARN_FILL

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    path = Path(path)
    wb.save(path)
    return path


def card_to_row(file_name: str, data) -> dict:
    row = {k: getattr(data, k) for k in data.field_names()}
    row["file"] = file_name
    return row

